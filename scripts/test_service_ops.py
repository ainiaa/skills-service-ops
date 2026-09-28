import io
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import time
import types
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import datetime
from pathlib import Path
from unittest.mock import patch


sys.path.insert(0, str(Path(__file__).parent))
import db_apply  # noqa: E402
import db_export  # noqa: E402
import db_query  # noqa: E402
import eval_contract  # noqa: E402
import log_analyzer  # noqa: E402
import paas  # noqa: E402
import setup  # noqa: E402
import sls_log_fetcher  # noqa: E402
import sls_query  # noqa: E402


class ReadOnlyQueryTests(unittest.TestCase):
    def test_only_allows_single_non_locking_select(self):
        self.assertTrue(paas.is_read_only_select("SELECT id FROM orders LIMIT 10;"))
        self.assertTrue(paas.is_read_only_select("SELECT 'SLEEP' AS note LIMIT 1;"))
        self.assertTrue(paas.is_read_only_select("SELECT id AS sleep FROM orders LIMIT 1;"))
        for sql in (
            "SELECT id FROM orders; DELETE FROM orders;",
            "SELECT id FROM orders LIMIT 1--2; DELETE FROM orders;",
            "SELECT id FROM orders LIMIT 1--\u00a02; DELETE FROM orders;",
            "SELECT id FROM orders FOR UPDATE",
            "SELECT id INTO OUTFILE '/tmp/orders' FROM orders",
            "SELECT id FROM orders LIMIT 1 /*!50000 FOR UPDATE */;",
            "SELECT id FROM orders LIMIT 1 /*!50000 INTO OUTFILE '/tmp/orders' */;",
            "SELECT GET_LOCK('service-ops-audit', 10) AS locked LIMIT 1;",
            "SELECT RELEASE_LOCK('service-ops-audit') AS released LIMIT 1;",
            "SELECT SLEEP(1) AS delayed LIMIT 1;",
            "SELECT BENCHMARK(1000000, SHA2('x', 256)) AS delayed LIMIT 1;",
            "SELECT LOAD_FILE('/etc/hosts') AS content LIMIT 1;",
            "SELECT `SLEEP`(1) AS delayed LIMIT 1;",
            "SELECT `GET_LOCK`('service-ops-audit', 10) AS locked LIMIT 1;",
            "SELECT `LOAD_FILE`('/etc/hosts') AS content LIMIT 1;",
            "SELECT LAST_INSERT_ID(1) AS session_state LIMIT 1;",
            "SELECT 1 INTO @service_ops_audit LIMIT 1;",
            "SELECT @service_ops_audit := 1 AS session_state LIMIT 1;",
        ):
            with self.subTest(sql=sql):
                self.assertFalse(paas.is_read_only_select(sql))

    def test_query_and_export_reject_quoted_dangerous_functions(self):
        sql = "SELECT `SLEEP`(1) AS delayed LIMIT 1"
        self.assertFalse(db_query.is_incident_query(sql))
        self.assertFalse(db_export.is_export_query(sql))

    def test_query_and_export_reject_unverified_sql_functions(self):
        for sql in (
            "SELECT mutate_orders(id) AS result FROM orders LIMIT 1",
            "SELECT app.mutate_orders(id) AS result FROM orders LIMIT 1",
            "SELECT `mutate_orders`(id) AS result FROM orders LIMIT 1",
            "SELECT `mutate``orders`(id) AS result FROM orders LIMIT 1",
            "SELECT mutate$COUNT(id) AS result FROM orders LIMIT 1",
            "SELECT 写入(id) AS result FROM orders LIMIT 1",
        ):
            with self.subTest(sql=sql):
                self.assertFalse(db_query.is_incident_query(sql))
                self.assertFalse(db_export.is_export_query(sql))
        self.assertTrue(db_query.is_incident_query("SELECT COUNT(id) AS total FROM orders LIMIT 1"))
        self.assertTrue(db_query.is_incident_query(
            "SELECT (price * quantity) AS total FROM orders WHERE (id = 1) LIMIT 1"))
        for function in ("IF(status=1, 1, 0)", "LEFT(name, 3)", "RIGHT(name, 3)",
                         "REPLACE(name, 'a', 'b')"):
            with self.subTest(function=function):
                self.assertTrue(db_query.is_incident_query(
                    "SELECT {} AS value FROM orders LIMIT 1".format(function)))

    def test_export_requires_a_numeric_limit(self):
        self.assertTrue(db_export.is_export_query("SELECT id FROM orders LIMIT 10"))
        self.assertFalse(db_export.is_export_query("SELECT id FROM orders"))

    def test_incident_query_requires_named_columns_and_a_small_limit(self):
        self.assertTrue(db_query.is_incident_query("SELECT id, status FROM orders LIMIT 10"))
        self.assertTrue(db_query.is_incident_query("SELECT price * quantity AS total FROM orders LIMIT 10"))
        for sql in (
            "SELECT * FROM orders LIMIT 10",
            "SELECT DISTINCT * FROM orders LIMIT 10",
            "SELECT HIGH_PRIORITY * FROM orders LIMIT 10",
            "SELECT DISTINCT HIGH_PRIORITY SQL_SMALL_RESULT * FROM orders LIMIT 10",
            "SELECT DISTINCT orders.* FROM orders LIMIT 10",
            "SELECT `orders`.* FROM orders LIMIT 10",
            "SELECT db.orders.* FROM db.orders LIMIT 10",
            "SELECT `x FROM y`, * FROM orders LIMIT 10",
            "SELECT id FROM orders",
            "SELECT id FROM orders LIMIT 101",
        ):
            with self.subTest(sql=sql):
                self.assertFalse(db_query.is_incident_query(sql))

    def test_incident_query_rejects_wildcards_in_later_select_clauses(self):
        for sql in (
            "SELECT id, name FROM users UNION SELECT * FROM users LIMIT 1",
            "SELECT id FROM users WHERE EXISTS (SELECT * FROM audit) LIMIT 1",
            "SELECT * LIMIT 1",
            "SELECT (SELECT id FROM users LIMIT 1) AS x, users.* FROM users LIMIT 1",
        ):
            with self.subTest(sql=sql):
                self.assertFalse(db_query.is_incident_query(sql))
        self.assertTrue(db_query.is_incident_query("SELECT COUNT(*) AS total FROM users LIMIT 1"))

    def test_read_only_join_using_is_not_treated_as_a_function(self):
        sql = "SELECT a.id FROM a JOIN b USING (id) LIMIT 1"
        self.assertTrue(db_query.is_incident_query(sql))
        self.assertTrue(db_export.is_export_query(sql))


class ExportTests(unittest.TestCase):
    def test_prod_query_reports_paas_error_as_failure(self):
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}}
        with patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "query", return_value={"error": "PaaS 查询请求失败。"}), \
             patch.object(sys, "argv", ["db_query.py", "--env", "prod", "--service", "orders",
                                         "--purpose", "核对订单", "--sql", "SELECT id FROM orders LIMIT 1"]), \
             redirect_stdout(io.StringIO()) as stdout, redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_query.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("PaaS 查询请求失败。", stderr.getvalue())

    def test_database_commands_require_a_purpose(self):
        commands = (
            (db_query.main, ["db_query.py", "--env", "prod", "--service", "orders",
                             "--sql", "SELECT id FROM orders LIMIT 1"]),
            (db_export.main, ["db_export.py", "--env", "prod", "--service", "orders",
                              "--sql", "SELECT id FROM orders LIMIT 1", "--output", "/tmp/orders.xlsx"]),
            (db_query.main, ["db_query.py", "--env", "prod", "--service", "orders",
                             "--purpose", "  ", "--sql", "SELECT id FROM orders LIMIT 1"]),
            (db_export.main, ["db_export.py", "--env", "prod", "--service", "orders",
                              "--purpose", "  ", "--sql", "SELECT id FROM orders LIMIT 1", "--output", "/tmp/orders.xlsx"]),
        )
        for command, argv in commands:
            with self.subTest(command=command.__module__), \
                 patch.object(sys, "argv", argv), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                command()
            self.assertEqual(exited.exception.code, 2)
            self.assertIn("--purpose", stderr.getvalue())

    def test_writes_xlsx_with_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "orders.xlsx"
            self.assertEqual(db_export.write_xlsx(["id"], [[1], [2]], output), 2)
            self.assertTrue(output.exists())

    def test_export_accepts_sql_file_or_stdin_and_rejects_oversized_results(self):
        self.assertEqual(db_export.read_sql("SELECT 1", None), "SELECT 1")
        self.assertTrue(db_export.is_export_query("SELECT id FROM orders LIMIT 10000"))
        self.assertFalse(db_export.is_export_query("SELECT id FROM orders LIMIT 10001"))

    def test_export_rejects_a_result_that_exceeds_the_approved_limit(self):
        with self.assertRaisesRegex(ValueError, "超过"):
            db_export.validate_export_result(["id"], [[1], [2]], 1)

    def test_export_rejects_non_tabular_rows(self):
        with self.assertRaisesRegex(ValueError, "行结构"):
            db_export.validate_export_result(["id"], ["not-a-row"], 1)

    def test_export_passes_the_approved_sql_limit_to_paas(self):
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}}
        sql = "SELECT id FROM orders LIMIT 10000"
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "query", return_value={"columns": ["id"], "rows": []}) as query, \
             patch.object(sys, "argv", ["db_export.py", "--env", "prod", "--service", "orders",
                                         "--purpose", "incident verification", "--sql", sql,
                                         "--output", str(Path(directory) / "orders.xlsx")]), \
             redirect_stdout(io.StringIO()) as stdout:
            db_export.main()
        query.assert_called_once_with(config, "orders", "prod", sql, 10000)
        self.assertEqual(json.loads(stdout.getvalue())["purpose"], "incident verification")

    def test_export_requires_an_absolute_xlsx_path(self):
        self.assertTrue(db_export.is_valid_output_path(Path("/tmp/orders.xlsx")))
        self.assertFalse(db_export.is_valid_output_path(Path("orders.xlsx")))
        self.assertFalse(db_export.is_valid_output_path(Path("/tmp/orders.csv")))

    def test_export_escapes_excel_formula_cells(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "orders.xlsx"
            db_export.write_xlsx(["=header"], [["=SUM(1,1)"], ["+cmd"], ["-1"], ["@value"], ["plain"]], output)
            sheet = load_workbook(output).active
        self.assertEqual([sheet.cell(row, 1).data_type for row in range(1, 6)], ["s"] * 5)
        self.assertEqual([sheet.cell(row, 1).value for row in range(1, 6)],
                         ["'=header", "'=SUM(1,1)", "'+cmd", "'-1", "'@value"])

    def test_export_reports_output_write_failure_without_traceback(self):
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}}
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "query", return_value={"columns": ["id"], "rows": [[1]]}), \
             patch.object(db_export, "write_xlsx", side_effect=OSError("permission denied")), \
             patch.object(sys, "argv", ["db_export.py", "--env", "prod", "--service", "orders",
                                         "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1",
                                         "--output", str(Path(directory) / "orders.xlsx")]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_export.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("Excel 文件写入失败。", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())

    def test_export_rejects_multiple_sql_sources(self):
        with patch.object(sys, "argv", ["db_export.py", "--env", "test", "--service", "orders",
                                         "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1", "--sql-file", "/tmp/orders.sql",
                                         "--output", "/tmp/orders.xlsx"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_export.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("not allowed with argument", stderr.getvalue())

    def test_export_reports_unreadable_sql_file_without_traceback(self):
        with patch.object(sys, "argv", ["db_export.py", "--env", "test", "--service", "orders",
                                         "--purpose", "incident verification", "--sql-file", "/definitely-missing-orders.sql", "--output", "/tmp/orders.xlsx"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_export.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("读取 SQL 文件失败。", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())


class ApplyTests(unittest.TestCase):
    def test_mixed_apply_tickets_preserve_sql_order(self):
        sql = ("INSERT INTO orders(id) VALUES (1); "
               "ALTER TABLE orders ADD COLUMN note VARCHAR(10); "
               "UPDATE orders SET note = 'ok' WHERE id = 1;")
        contract = {"service_name": "service-ops", "fields": {
            "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs"}}
        tickets = db_apply._tickets(db_apply.plan_changes(sql), "orders-db", "repair", "production", "test", 2000, contract)
        self.assertEqual([ticket["kind"] for ticket in tickets], ["DML", "DDL", "DML"])
        self.assertEqual([ticket["statements"][0].split()[0] for ticket in tickets], ["INSERT", "ALTER", "UPDATE"])

    def test_apply_rejects_non_mutating_sql(self):
        with self.assertRaises(ValueError):
            db_apply.plan_changes("SELECT id FROM orders;")

    def test_apply_rejects_unbounded_dml(self):
        for sql in ("UPDATE orders SET state = 1;", "DELETE FROM orders;"):
            with self.subTest(sql=sql):
                with self.assertRaises(ValueError):
                    db_apply.plan_changes(sql)

    def test_apply_rejects_arithmetic_double_dash_that_hides_a_second_statement(self):
        with self.assertRaises(ValueError):
            db_apply.plan_changes("UPDATE orders SET state = 1 WHERE id = 1--2; DELETE FROM orders;")

    def test_apply_rejects_mysql_executable_comments(self):
        with self.assertRaises(ValueError):
            db_apply.plan_changes("UPDATE orders SET state = 1 /*!50000 WHERE id = 1 */;")

    def test_apply_rejects_where_that_only_exists_in_a_nested_subquery(self):
        with self.assertRaisesRegex(ValueError, "WHERE"):
            db_apply.plan_changes("UPDATE orders SET status = (SELECT status FROM source WHERE id = 1);")

    def test_apply_accepts_a_semicolon_inside_a_sql_comment(self):
        plan = db_apply.plan_changes("/* migration; note */ UPDATE orders SET status = 1 WHERE id = 1;")
        self.assertEqual(plan["dml"], ["/* migration; note */ UPDATE orders SET status = 1 WHERE id = 1;"])

    def test_apply_requires_explicit_opt_in_for_destructive_ddl(self):
        with self.assertRaises(ValueError):
            db_apply.plan_changes("TRUNCATE TABLE orders;")
        self.assertEqual(db_apply.plan_changes("TRUNCATE TABLE orders;", allow_destructive=True)["ddl"],
                         ["TRUNCATE TABLE orders;"])

    def test_confirmation_token_binds_exact_sql(self):
        first = db_apply.confirmation_token("prod", "wms", "group", "repair", ["UPDATE t SET n = 1;"])
        second = db_apply.confirmation_token("prod", "wms", "group", "repair", ["UPDATE t SET n = 2;"])
        self.assertTrue(first.startswith("apply-"))
        self.assertNotEqual(first, second)

    def test_cleanup_removes_only_expired_submitted_preflights(self):
        now = datetime(2026, 9, 4, 12, 0, 0)
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"):
            directory_path = paas.CONFIG_PATH.parent / "preflights"
            directory_path.mkdir()
            expired = directory_path / ("apply-" + "a" * 32 + ".submitted")
            recent = directory_path / ("apply-" + "b" * 32 + ".submitted")
            pending = directory_path / ("apply-" + "c" * 32 + ".json")
            for path in (expired, recent, pending):
                path.write_text("{}", encoding="utf-8")
            os.utime(expired, (now.timestamp() - 8 * 24 * 60 * 60,) * 2)
            os.utime(recent, (now.timestamp() - 6 * 24 * 60 * 60,) * 2)
            self.assertEqual(db_apply.cleanup_submitted_preflights(now=now), 1)
            self.assertFalse(expired.exists())
            self.assertTrue(recent.exists())
            self.assertTrue(pending.exists())

    def test_save_preflight_cleans_expired_submitted_preflights(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"):
            directory_path = paas.CONFIG_PATH.parent / "preflights"
            directory_path.mkdir()
            expired = directory_path / ("apply-" + "a" * 32 + ".submitted")
            expired.write_text("{}", encoding="utf-8")
            os.utime(expired, (0, 0))
            db_apply.save_preflight("apply-" + "d" * 32, {"profile": "default"}, [])
            self.assertFalse(expired.exists())

    def test_claim_preflight_redacts_consumed_plan(self):
        token = "apply-" + "a" * 32
        request = {"profile": "default", "reason": "repair"}
        tickets = [{"body": {"sql": "UPDATE orders SET status = 1 WHERE id = 1;"}}]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"):
            db_apply.save_preflight(token, request, tickets, "nonce")
            self.assertEqual(db_apply.claim_preflight(token)["tickets"], tickets)
            consumed = json.loads(db_apply.submitted_preflight_path(token).read_text(encoding="utf-8"))
        self.assertIn("consumed_at", consumed)
        self.assertNotIn("request", consumed)
        self.assertNotIn("tickets", consumed)

    def test_submit_rejects_a_different_ticket_split_than_the_preflight(self):
        class Session:
            def post(self, *_args, **_kwargs):
                return type("Response", (), {"raise_for_status": lambda self: None,
                                            "json": lambda self: {"id": 1}})()

        sql = "UPDATE orders SET status = 1 WHERE id = 1; UPDATE orders SET status = 2 WHERE id = 2;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"test": "testing"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "service-ops", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
            }},
        }}
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml") as config_path, \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()), \
             patch.object(sys, "argv", ["db_apply.py", "--env", "test", "--service", "orders",
                                         "--reason", "repair", "--batch-size", "1", "--sql", sql]), \
             redirect_stdout(io.StringIO()) as output:
            db_apply.main()
            token = json.loads(output.getvalue())["confirmation_token"]
            stored = db_apply.load_preflight(token)
        with tempfile.TemporaryDirectory() as submit_directory, \
             patch.object(paas, "CONFIG_PATH", Path(submit_directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()), \
             patch.object(sys, "argv", ["db_apply.py", "--env", "test", "--service", "orders",
                                         "--reason", "repair", "--batch-size", "2", "--sql", sql,
                                         "--submit", "--confirm", token]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_apply.save_preflight(token, stored["request"], stored["tickets"], stored.get("nonce"))
            db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("预检计划", stderr.getvalue())

    def test_submit_uses_the_preflight_ticket_body_without_rebuilding_it(self):
        submitted = []

        class Session:
            def post(self, _url, **kwargs):
                submitted.append(kwargs["json"])
                return type("Response", (), {"raise_for_status": lambda self: None,
                                            "json": lambda self: {"id": 1}})()

        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"prod": "production"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "service-ops", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
                "start": "start", "end": "end",
            }},
        }}
        arguments = ["db_apply.py", "--env", "prod", "--service", "orders", "--reason", "repair", "--sql", sql]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml") as config_path, \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()), \
             patch.object(sys, "argv", arguments), \
             redirect_stdout(io.StringIO()) as output:
            db_apply.main()
            preview = json.loads(output.getvalue())
            stored = db_apply.load_preflight(preview["confirmation_token"])
        with tempfile.TemporaryDirectory() as submit_directory, \
             patch.object(paas, "CONFIG_PATH", Path(submit_directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()), \
             patch.object(db_apply, "_tickets", side_effect=AssertionError("must use the stored preflight plan")), \
             patch.object(sys, "argv", arguments + ["--submit", "--confirm", preview["confirmation_token"]]), \
             redirect_stdout(io.StringIO()):
            db_apply.save_preflight(preview["confirmation_token"], stored["request"], stored["tickets"], stored.get("nonce"))
            db_apply.main()
        self.assertEqual(submitted, [preview["tickets"][0]["body"]])

    def test_submit_rejects_a_preflight_from_a_different_profile(self):
        submitted = []

        class Session:
            def post(self, *_args, **_kwargs):
                submitted.append(True)
                return type("Response", (), {"raise_for_status": lambda self: None,
                                            "json": lambda self: {"id": 1}})()

        def config_for_profile():
            profile = paas.active_profile()
            return {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
                "apply_envs": {"test": profile + "-test"}, "apply_api_base": "https://" + profile + ".invalid",
                "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
                "apply_contract": {"service_name": profile + "-service", "fields": {
                    "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
                }},
            }}

        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        base = ["db_apply.py", "--env", "test", "--service", "orders", "--reason", "repair", "--sql", sql]
        with tempfile.TemporaryDirectory() as directory, \
             patch.dict(os.environ, {"SERVICE_OPS_PROFILE": "default"}, clear=False), \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "load_config", side_effect=config_for_profile), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            with patch.object(sys, "argv", ["db_apply.py", "--profile", "regional", *base[1:]]), \
                 redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            with patch.object(sys, "argv", ["db_apply.py", "--profile", "default", *base[1:], "--submit", "--confirm", token]), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("预检计划不一致", stderr.getvalue())
        self.assertEqual(submitted, [])

    def test_submit_rejects_a_preflight_when_the_paas_contract_changes(self):
        submitted = []

        class Session:
            def post(self, *_args, **_kwargs):
                submitted.append(True)
                return type("Response", (), {"raise_for_status": lambda self: None,
                                            "json": lambda self: {"id": 1}})()

        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"test": "testing"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "first-contract", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
            }},
        }}
        base = ["db_apply.py", "--env", "test", "--service", "orders", "--reason", "repair", "--sql", sql]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            with patch.object(sys, "argv", base), redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            config["paas"]["apply_contract"]["service_name"] = "second-contract"
            with patch.object(sys, "argv", base + ["--submit", "--confirm", token]), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("预检计划不一致", stderr.getvalue())
        self.assertEqual(submitted, [])

    def test_submit_rejects_a_preflight_when_the_apply_environment_changes(self):
        submitted = []

        class Session:
            def post(self, *_args, **_kwargs):
                submitted.append(True)
                return type("Response", (), {"raise_for_status": lambda self: None,
                                            "json": lambda self: {"id": 1}})()

        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"test": "testing"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "service-ops", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
            }},
        }}
        base = ["db_apply.py", "--env", "test", "--service", "orders", "--reason", "repair", "--sql", sql]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            with patch.object(sys, "argv", base), redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            config["paas"]["apply_envs"]["test"] = "different-testing"
            with patch.object(sys, "argv", base + ["--submit", "--confirm", token]), \
                 redirect_stdout(io.StringIO()), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                db_apply.main()
            self.assertEqual(db_apply.load_preflight(token)["tickets"][0]["body"]["envs"], ["testing"])
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("预检计划不一致", stderr.getvalue())
        self.assertEqual(submitted, [])

    def test_submit_consumes_a_confirmation_token_after_the_first_post(self):
        calls = []

        class Session:
            def post(self, *_args, **_kwargs):
                calls.append(True)
                return type("Response", (), {"raise_for_status": lambda self: None,
                                            "json": lambda self: {"id": len(calls)}})()

        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"test": "testing"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "service-ops", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
            }},
        }}
        base = ["db_apply.py", "--env", "test", "--service", "orders", "--reason", "repair", "--sql", sql]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            with patch.object(sys, "argv", base), redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            with patch.object(sys, "argv", base + ["--submit", "--confirm", token]), redirect_stdout(io.StringIO()):
                db_apply.main()
            with patch.object(sys, "argv", base + ["--submit", "--confirm", token]), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("已提交", stderr.getvalue())
        self.assertEqual(len(calls), 1)

    def test_apply_submit_keeps_preflight_when_cookie_is_missing(self):
        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"test": "testing"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "service-ops", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
            }},
        }}
        base = ["db_apply.py", "--env", "test", "--service", "orders", "--reason", "repair", "--sql", sql]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value=""):
            with patch.object(sys, "argv", base), redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            with patch.object(sys, "argv", base + ["--submit", "--confirm", token]), \
                 redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                db_apply.main()
            self.assertEqual(db_apply.load_preflight(token)["request"]["service"], "orders")

    def test_schedule_is_only_sent_for_production(self):
        contract = {"service_name": "service-ops", "fields": {
            "sql": "sql", "service": "service", "group": "group", "reason": "reason",
            "envs": "environments", "start": "starts_at", "end": "ends_at",
        }}
        prod = db_apply.build_apply_body(["UPDATE t SET n = 1;"], "group", "repair", "prod", "2026-09-01 10:00:00", "2026-09-01 12:00:00", contract)
        test = db_apply.build_apply_body(["UPDATE t SET n = 1;"], "group", "repair", "test", None, None, contract)
        self.assertEqual(prod["service"], "service-ops")
        self.assertIn("starts_at", prod)
        self.assertNotIn("starts_at", test)

    def test_task_actions_build_the_configured_contract(self):
        actions = {
            "approve": {"endpoint": "/tasks/approve", "id_field": "task_id", "message_field": "note"},
            "execute": {"endpoint": "/tasks/run", "id_field": "task_id"},
            "recall": {"endpoint": "/tasks/cancel", "id_field": "task_id"},
        }
        url, body = db_apply.build_task_op("approve", 7, actions, "ok")
        self.assertEqual(url, "/tasks/approve")
        self.assertEqual(body, {"task_id": 7, "note": "ok"})
        self.assertEqual(db_apply.build_task_op("execute", 8, actions, "x")[1], {"task_id": 8})
        self.assertEqual(db_apply.build_task_op("recall", 9, actions, "x")[1], {"task_id": 9})

    def test_task_action_rejects_a_preflight_from_a_different_profile(self):
        submitted = []

        class Session:
            def post(self, *_args, **_kwargs):
                submitted.append(True)
                return type("Response", (), {"raise_for_status": lambda self: None})()

        def config_for_profile():
            return {"paas": {"task_api_base": "https://shared.invalid", "task_actions": {
                "execute": {"endpoint": "/tasks/run", "id_field": "task_id"},
            }}}

        base = ["db_apply.py", "--action", "execute", "--id", "7"]
        with tempfile.TemporaryDirectory() as directory, \
             patch.dict(os.environ, {"SERVICE_OPS_PROFILE": "default"}, clear=False), \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "load_config", side_effect=config_for_profile), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            with patch.object(sys, "argv", ["db_apply.py", "--profile", "regional", *base[1:]]), \
                 redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            with patch.object(sys, "argv", ["db_apply.py", "--profile", "default", *base[1:], "--confirm", token]), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("预检计划不一致", stderr.getvalue())
        self.assertEqual(submitted, [])

    def test_task_action_consumes_a_confirmation_token_after_the_first_post(self):
        calls = []

        class Session:
            def post(self, *_args, **_kwargs):
                calls.append(True)
                return type("Response", (), {"raise_for_status": lambda self: None})()

        config = {"paas": {"task_api_base": "https://example.invalid", "task_actions": {
            "execute": {"endpoint": "/tasks/run", "id_field": "task_id"},
        }}}
        base = ["db_apply.py", "--action", "execute", "--id", "7"]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            with patch.object(sys, "argv", base), redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            with patch.object(sys, "argv", base + ["--confirm", token]), redirect_stdout(io.StringIO()):
                db_apply.main()
            with patch.object(sys, "argv", base + ["--confirm", token]), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("已提交", stderr.getvalue())
        self.assertEqual(len(calls), 1)

    def test_task_action_keeps_preflight_when_cookie_is_missing(self):
        config = {"paas": {"task_api_base": "https://example.invalid", "task_actions": {
            "execute": {"endpoint": "/tasks/run", "id_field": "task_id"},
        }}}
        base = ["db_apply.py", "--action", "execute", "--id", "7"]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value=""):
            with patch.object(sys, "argv", base), redirect_stdout(io.StringIO()) as output:
                db_apply.main()
                token = json.loads(output.getvalue())["confirmation_token"]
            with patch.object(sys, "argv", base + ["--confirm", token]), \
                 redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                db_apply.main()
            self.assertEqual(db_apply.load_preflight(token)["request"]["ids"], [7])

    def test_apply_contract_requires_configured_fields_and_endpoints(self):
        with self.assertRaises(ValueError):
            db_apply.build_apply_body(["UPDATE t SET n = 1;"], "group", "repair", "prod", None, None, {})
        with self.assertRaises(ValueError):
            db_apply.build_task_op("approve", 1, {}, "ok")

    def test_apply_response_uses_the_configured_id_field(self):
        self.assertEqual(db_apply.apply_id({"change_id": 12}, "change_id"), "12")
        with self.assertRaises(RuntimeError):
            db_apply.apply_id({}, "change_id")

    def test_apply_preflight_reports_incomplete_contract_without_traceback(self):
        config = {"services": {"wms": {"prod_db_group": "wms-db"}},
                  "paas": {"apply_envs": {"prod": "production"}, "apply_contract": {}}}
        with patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(sys, "argv", ["db_apply.py", "--env", "prod", "--service", "wms", "--reason", "repair", "--sql", "UPDATE orders SET status = 1 WHERE id = 1;"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("PaaS apply_contract 未配置完整字段。", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())

    def test_apply_submit_reports_a_network_failure_without_traceback(self):
        class RequestException(Exception):
            pass

        class FailingSession:
            def post(self, *_args, **_kwargs):
                raise RequestException("offline")

        sql = "UPDATE orders SET status = 1 WHERE id = 1;"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "apply_envs": {"prod": "production"}, "apply_api_base": "https://example.invalid",
            "apply_endpoints": {"ddl": "ddl", "dml": "dml"}, "apply_response_id_field": "id",
            "apply_contract": {"service_name": "service-ops", "fields": {
                "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
                "start": "start", "end": "end",
            }},
        }}
        plan = db_apply.plan_changes(sql)
        request = {"profile": paas.active_profile(), "env": "prod", "service": "orders", "group": "orders-db",
                   "reason": "repair", "statements": ["DML:" + sql], "batch_size": 2000,
                   "target": {"api_base": "https://example.invalid", "endpoints": {"ddl": "ddl", "dml": "dml"},
                              "response_id_field": "id", "environment": "production",
                              "contract": config["paas"]["apply_contract"]}}
        tickets = db_apply._tickets(plan, "orders-db", "repair", "production", "prod", 2000,
                                    config["paas"]["apply_contract"])
        token = db_apply.confirmation_token("prod", "orders", "orders-db", "repair",
                                            {"request": request, "tickets": tickets})
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(paas, "CONFIG_PATH", Path(directory) / "config.yaml"), \
             patch.dict(sys.modules, {"requests": types.SimpleNamespace(RequestException=RequestException)}), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=FailingSession()), \
             patch.object(sys, "argv", ["db_apply.py", "--env", "prod", "--service", "orders",
                                         "--reason", "repair", "--sql", sql, "--submit", "--confirm", token]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_apply.save_preflight(token, request, tickets)
            db_apply.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("PaaS 工单提交失败。", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())


class DatabaseRoutingTests(unittest.TestCase):
    def test_flat_configuration_does_not_reuse_default_routing_for_another_profile(self):
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "config.yaml"
            config_path.write_text("paas:\n  query_api_url: https://default.invalid/query\n", encoding="utf-8")
            config_path.chmod(0o600)
            with patch.object(paas, "CONFIG_PATH", config_path):
                self.assertEqual(paas.load_config("default")["paas"]["query_api_url"],
                                 "https://default.invalid/query")
                with self.assertRaisesRegex(ValueError, "未配置 profile：regional"):
                    paas.load_config("regional")
                config_path.write_text("profiles: {}\n", encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "未配置 profile：default"):
                    paas.load_config("default")
                config_path.write_text("profiles:\n  regional:\n    paas:\n"
                                       "      query_api_url: https://regional.invalid/query\n", encoding="utf-8")
                self.assertEqual(paas.load_config("regional")["paas"]["query_api_url"],
                                 "https://regional.invalid/query")

    def test_prod_query_uses_the_sql_limit_as_its_paas_result_cap(self):
        sql = "SELECT id FROM orders LIMIT 7"
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}}
        with patch.object(paas, "query", return_value={"columns": ["id"], "rows": []}) as query:
            db_query.query_prod(sql, "orders", config)
        query.assert_called_once_with(config, "orders", "prod", sql, 7)

    def test_prod_query_rejects_a_paas_response_beyond_the_sql_limit(self):
        class Response:
            def raise_for_status(self):
                pass

            def json(self):
                return [{"columnList": ["id"], "rows": [[index] for index in range(8)]}]

        class Session:
            def post(self, _url, **kwargs):
                self.payload = kwargs["json"]
                return Response()

        session = Session()
        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "query_api_url": "https://example.invalid/query", "query_envs": {"prod": "production"},
            "query_contract": {"fields": {"group": "group", "env": "env", "sql": "sql", "limit": "limit"}},
        }}
        with patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=session):
            result = db_query.query_prod("SELECT id FROM orders LIMIT 7", "orders", config)
        self.assertEqual(session.payload["limit"], 7)
        self.assertEqual(result, {"error": "PaaS 查询返回结构无效。"})

    def test_runtime_profile_does_not_come_from_environment(self):
        with patch.dict(os.environ, {"SERVICE_OPS_PROFILE": "regional"}, clear=False):
            paas.set_active_profile(None)
            self.assertEqual(paas.active_profile(), "default")

    def test_sls_connection_does_not_fall_back_to_environment_configuration(self):
        with patch.dict(os.environ, {
            "SLS_LOG_PROJECT": "environment-project",
            "SLS_LOG_REGION": "ap-southeast-1",
            "SLS_LOG_ENDPOINT": "environment-endpoint",
        }, clear=False):
            with self.assertRaisesRegex(ValueError, "SLS project"):
                sls_query.resolve_connection("default", None, None, None, {"sls": {}})
        self.assertEqual(
            sls_query.resolve_connection("default", None, None, None,
                                         {"sls": {"project": "configured-project", "region": "cn-hangzhou"}}),
            ("configured-project", "cn-hangzhou", None),
        )

    def test_keychain_timeout_is_bounded_and_falls_back_to_private_config(self):
        with patch.object(paas.subprocess, "run", side_effect=subprocess.TimeoutExpired("security", 10)) as run, \
             patch.dict(os.environ, {"PAAS_COOKIE": "environment-cookie"}, clear=False):
            self.assertEqual(paas.get_cookie(config={"paas": {"cookie": "configured-cookie"}}), "configured-cookie")
        self.assertEqual(run.call_args.kwargs["timeout"], paas.KEYCHAIN_TIMEOUT_SECONDS)

    def test_get_cookie_uses_the_requested_profile_keychain_entry(self):
        with patch.object(paas, "_keychain_value", return_value="cookie") as keychain_value:
            self.assertEqual(paas.get_cookie("regional", {"paas": {"cookie": "configured-cookie"}}), "cookie")
        keychain_value.assert_called_once_with("paas-cookie", "regional")

    def test_test_database_keychain_password_takes_precedence_over_private_config(self):
        with patch.object(paas, "_keychain_value", return_value="keychain-password"):
            self.assertEqual(paas.get_test_db_password("regional", {"test_db": {"password": "configured-password"}}),
                             "keychain-password")

    def test_paas_credentials_fall_back_to_the_resolved_private_config(self):
        with patch.object(paas, "_keychain_value", return_value=""), \
             patch.dict(os.environ, {"PAAS_COOKIE": "environment-cookie", "TEST_DB_PASSWORD": "environment-password"}, clear=False):
            self.assertEqual(paas.get_cookie(config={"paas": {"cookie": "configured-cookie"}}), "configured-cookie")
            self.assertEqual(paas.get_test_db_password(config={"test_db": {"password": "configured-password"}}),
                             "configured-password")
        with patch.object(paas, "_keychain_value", return_value=""), \
             patch.dict(os.environ, {"PAAS_COOKIE": "environment-cookie"}, clear=False):
            self.assertEqual(paas.get_cookie(config={"paas": {}}), "")

    def test_paas_query_uses_the_caller_limit(self):
        class RequestException(Exception):
            pass

        class Response:
            def raise_for_status(self):
                return None

            def json(self):
                return [{"columnList": ["id"], "rows": []}]

        captured = {}

        class Session:
            def post(self, _url, **kwargs):
                captured.update(kwargs)
                return Response()

        config = {"services": {"orders": {"prod_db_group": "orders-db"}}, "paas": {
            "query_api_url": "https://example.invalid/query", "query_envs": {"prod": "production"},
            "query_contract": {"fields": {"group": "group", "env": "env", "sql": "sql", "limit": "limit"}},
        }}
        with patch.dict(sys.modules, {"requests": types.SimpleNamespace(RequestException=RequestException)}), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=Session()):
            self.assertEqual(paas.query(config, "orders", "prod", "SELECT id FROM orders LIMIT 10000", 10000)["rows"], [])
        self.assertEqual(captured["json"]["limit"], 10000)

    def test_test_database_route_is_available(self):
        self.assertTrue(callable(db_query.query_test))
        self.assertTrue(callable(db_query.query_prod))

    def test_test_database_query_requires_a_configured_database(self):
        with self.assertRaisesRegex(ValueError, "测试数据库"):
            db_query.query_test("SELECT id FROM orders LIMIT 1", "", {"test_db": {}})

    def test_test_database_mysql_error_is_reported_without_traceback(self):
        class MySQLError(Exception):
            pass

        fake_pymysql = types.SimpleNamespace(MySQLError=MySQLError,
                                             connect=lambda **_kwargs: (_ for _ in ()).throw(MySQLError("denied")))
        config = {"test_db": {"host": "db.example.invalid", "user": "readonly"},
                  "services": {"orders": {"test_database": "orders_test"}}}
        with patch.dict(sys.modules, {"pymysql": fake_pymysql}), \
             patch.object(paas, "set_active_profile"), \
             patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "_keychain_value", return_value="password"), \
             patch.object(sys, "argv", ["db_query.py", "--env", "test", "--service", "orders",
                                         "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            db_query.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("测试库查询失败。", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())

    def test_paas_headers_use_configured_cookie_rule_and_headers(self):
        config = {"paas": {"cookie_prefix": "Session=", "request_headers": {"origin": "https://example.invalid"}}}
        headers = paas.api_headers(config, "token")
        self.assertEqual(headers["cookie"], "Session=token")
        self.assertEqual(headers["origin"], "https://example.invalid")

    def test_paas_query_payload_uses_configured_field_names(self):
        config = {"paas": {"query_envs": {"prod": "production"}, "query_contract": {"fields": {
            "group": "database", "env": "environment", "sql": "statement", "limit": "max_rows",
        }}}, "services": {"orders": {"prod_db_group": "orders-db"}}}
        self.assertEqual(paas.query_payload(config, "orders", "prod", "SELECT id FROM orders", 100), {
            "database": "orders-db", "environment": "production", "statement": "SELECT id FROM orders", "max_rows": 100,
        })

    def test_paas_query_rejects_an_invalid_success_response_shape(self):
        config = {"paas": {"query_api_url": "https://example.invalid/query",
                           "query_envs": {"prod": "production"}, "query_contract": {"fields": {
                               "group": "database", "env": "environment", "sql": "statement", "limit": "max_rows",
                           }}}, "services": {"orders": {"prod_db_group": "orders-db"}}}
        response = type("Response", (), {"data": ["invalid"], "raise_for_status": lambda self: None,
                                           "json": lambda self: self.data})()
        session = type("Session", (), {"post": lambda self, *args, **kwargs: response})()
        with patch.dict(sys.modules, {"requests": types.SimpleNamespace(RequestException=Exception)}), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(paas, "build_session", return_value=session):
            self.assertEqual(paas.query(config, "orders", "prod", "SELECT id FROM orders LIMIT 1"),
                             {"error": "PaaS 查询返回结构无效。"})
            response.data = {"unexpected": "object"}
            self.assertEqual(paas.query(config, "orders", "prod", "SELECT id FROM orders LIMIT 1"),
                             {"error": "PaaS 查询返回结构无效。"})
            for invalid in ([], [{}], [{"columnList": ["id"]}, {"columnList": ["id"], "rows": [[1]]}],
                            [{"columnList": ["id"], "rows": []}, {"columnList": ["id"], "rows": [[1]]}],
                            [{"columnList": ["id"], "rows": [[1, 2]]}],
                            [{"columnList": [1], "rows": [[1]]}],
                            [{"columnList": ["id"], "rows": [[1], [2]]}]):
                with self.subTest(response=invalid):
                    response.data = invalid
                    self.assertEqual(paas.query(config, "orders", "prod", "SELECT id FROM orders LIMIT 1", limit=1),
                                     {"error": "PaaS 查询返回结构无效。"})

    def test_load_config_rejects_invalid_yaml_and_non_mapping_root(self):
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "config.yaml"
            config_path.write_text("profiles: [", encoding="utf-8")
            config_path.chmod(0o600)
            invalid_yaml = types.SimpleNamespace(YAMLError=RuntimeError,
                                                 safe_load=lambda _file: (_ for _ in ()).throw(RuntimeError()))
            with patch.dict(sys.modules, {"yaml": invalid_yaml}), \
                 patch.object(paas, "CONFIG_PATH", config_path), self.assertRaisesRegex(ValueError, "配置文件"):
                paas.load_config()
            config_path.write_text("- not-a-config", encoding="utf-8")
            mapping_root = types.SimpleNamespace(YAMLError=RuntimeError, safe_load=lambda _file: ["not-a-config"])
            with patch.dict(sys.modules, {"yaml": mapping_root}), \
                 patch.object(paas, "CONFIG_PATH", config_path), self.assertRaisesRegex(ValueError, "配置文件"):
                paas.load_config()

    def test_redacts_common_sensitive_log_fields(self):
        value = "Authorization: Bearer top-secret; password=hunter2; request=ok"
        redacted = paas.redact_text(value)
        self.assertNotIn("top-secret", redacted)
        self.assertNotIn("hunter2", redacted)
        self.assertIn("request=ok", redacted)

    def test_redacts_sensitive_log_mapping_values(self):
        record = paas.redact_log_fields({
            "authorization": "Bearer top-secret", "accessToken": "another-secret", "message": "token=third-secret", "count": 1,
        })
        self.assertEqual(record["authorization"], "***")
        self.assertEqual(record["accessToken"], "***")
        self.assertNotIn("third-secret", record["message"])
        self.assertEqual(record["count"], 1)

    def test_redacts_private_patterns_in_raw_log_values(self):
        rules = paas.log_analysis_rules({"log_analysis": {"redaction_patterns": [r"C-\d+"]}})
        record = paas.redact_log_fields({"message": "customer=C-42", "nested": ["C-43"]}, rules["redaction_patterns"])
        self.assertNotIn("C-42", json.dumps(record))
        self.assertNotIn("C-43", json.dumps(record))

    def test_log_analysis_rules_reject_invalid_patterns_and_shapes(self):
        for analysis in (
            {"redaction_patterns": ["("]},
            {"related_service_pattern": "("},
            {"entity_patterns": {"invoice": "("}},
            {"entity_patterns": []},
        ):
            with self.subTest(analysis=analysis), self.assertRaises(ValueError):
                paas.log_analysis_rules({"log_analysis": analysis})


class SetupTests(unittest.TestCase):
    def test_readme_includes_a_scenario_based_usage_guide(self):
        readme = (Path(__file__).parent.parent / "README.md").read_text(encoding="utf-8")
        for section in ("## 快速开始", "## 常用场景", "### 查询日志", "### TraceId 排障",
                        "### 数据验证与 Excel 导出", "### 工单预检与提交", "## 常见问题"):
            with self.subTest(section=section):
                self.assertIn(section, readme)

    def test_configuration_template_points_to_the_runtime_configuration_path(self):
        template = (Path(__file__).parent.parent / "config" / "settings.yaml.example").read_text(encoding="utf-8")
        self.assertIn("~/.service-ops/config.yaml", template)
        self.assertIn("access_key", template)
        self.assertIn("access_secret", template)
        self.assertNotIn("复制为 settings.yaml", template)

    def test_clean_install_workflow_uses_the_lock_and_readiness_check(self):
        workflow = (Path(__file__).parent.parent / ".github" / "workflows" / "verify.yml").read_text(encoding="utf-8")
        self.assertIn("runner: [macos-latest, ubuntu-latest]", workflow)
        self.assertIn("runs-on: ${{ matrix.runner }}", workflow)
        self.assertIn("actions/checkout@11bd71901bbe5b1630ceea73d27597364c9af683", workflow)
        self.assertIn("actions/setup-python@a26af69be951a213d495a4c3e4e4022e16d87065", workflow)
        self.assertIn("python -m pip install -r requirements.lock", workflow)
        self.assertIn("python scripts/setup.py --check-dependencies --capability all", workflow)

    def test_dependency_lock_pins_every_runtime_dependency(self):
        root = Path(__file__).parent.parent
        direct_dependencies = {"requests", "pyyaml", "aliyun-log-python-sdk", "openpyxl", "pymysql"}
        lock_lines = [line for line in (root / "requirements.lock").read_text(encoding="utf-8").splitlines()
                      if line and not line.startswith("#")]
        self.assertTrue(lock_lines)
        self.assertTrue(all("==" in line for line in lock_lines))
        self.assertTrue(direct_dependencies.issubset({line.split("==", 1)[0].lower() for line in lock_lines}))

    def test_check_readiness_uses_requested_profile_for_paas_credentials(self):
        config = {"paas": {"query_api_url": "https://example.invalid/query",
                           "query_envs": {"prod": "production", "test": "testing"},
                           "query_contract": {"fields": {
                               "group": "database", "env": "environment",
                               "sql": "statement", "limit": "max_rows",
                           }}}}
        with patch.object(paas, "load_config", return_value=config) as load_config, \
             patch.object(paas, "get_cookie", return_value="cookie") as get_cookie, \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("paas", "regional"), [])
        load_config.assert_called_once_with("regional")
        get_cookie.assert_called_once_with("regional", config)

    def test_check_readiness_reports_missing_configuration_and_credentials(self):
        with patch.object(paas, "load_config", side_effect=ValueError("配置缺失")), \
             patch.object(paas, "get_cookie", return_value=""), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness(), ["配置缺失", "PaaS Cookie 未配置"])

    def test_check_readiness_reports_missing_config_parser_without_traceback(self):
        with patch.object(paas, "load_config", side_effect=ImportError("缺少 pyyaml。")), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("db"), ["缺少 pyyaml。"])

    def test_check_readiness_does_not_load_config_when_pyyaml_is_missing(self):
        with patch.object(paas, "load_config") as load_config, \
             patch.object(setup, "missing_dependencies", return_value=["pyyaml"]):
            self.assertEqual(setup.check_readiness("db"), ["缺少 Python 依赖：pyyaml"])
        load_config.assert_not_called()

    def test_check_readiness_checks_sls_configuration_and_credentials(self):
        config = {"paas": {"query_api_url": "https://example.invalid/query",
                           "query_envs": {"prod": "production", "test": "testing"},
                           "query_contract": {"fields": {
                               "group": "database", "env": "environment",
                               "sql": "statement", "limit": "max_rows",
                           }}}, "sls": {}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(sls_query, "get_credentials", side_effect=ValueError("SLS AK/SK 未配置")), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(
                setup.check_readiness("sls"),
                ["当前 profile 未配置 SLS project。", "当前 profile 未配置 SLS region 或 endpoint。", "SLS AK/SK 未配置"],
            )

    def test_check_readiness_rejects_invalid_log_analysis_rules(self):
        config = {"sls": {"project": "project", "region": "region"},
                  "log_analysis": {"entity_patterns": {"invoice": "("}}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("sls"), ["log_analysis.entity_patterns.invoice 包含无效正则。"])

    def test_check_readiness_rejects_incomplete_paas_query_contract(self):
        with patch.object(paas, "load_config", return_value={"paas": {}}), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(sls_query, "get_credentials") as credentials, \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("paas"), [
                "PaaS 查询端点未配置。",
                "PaaS query_contract 未配置完整字段。",
                "PaaS 查询环境映射未配置：prod、test。",
            ])
        credentials.assert_not_called()

    def test_check_readiness_accepts_complete_paas_query_contract(self):
        config = {"paas": {"query_api_url": "https://example.invalid/query",
                           "query_envs": {"prod": "production", "test": "testing"},
                           "query_contract": {"fields": {
                               "group": "database", "env": "environment",
                               "sql": "statement", "limit": "max_rows",
                           }}}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(sls_query, "get_credentials") as credentials, \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("paas"), [])
        credentials.assert_not_called()

    def test_check_readiness_checks_database_and_apply_requirements(self):
        config = {"test_db": {}, "paas": {}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_cookie", return_value="cookie"), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("db"), ["测试库 host 未配置。", "测试库 user 未配置。"])
            self.assertEqual(setup.check_readiness("apply"), [
                "PaaS apply 提交端点或响应标识字段未配置。",
                "PaaS apply 环境映射未配置：prod、test。",
                "PaaS apply_contract 未配置完整字段。",
                "PaaS 工单操作端点或字段未配置。",
            ])

    def test_apply_readiness_checks_existing_ticket_actions(self):
        config = {"paas": {"apply_api_base": "https://example.invalid", "apply_response_id_field": "id",
                           "apply_endpoints": {"ddl": "ddl", "dml": "dml"},
                           "apply_envs": {"prod": "production", "test": "testing"},
                           "apply_contract": {"service_name": "service-ops", "fields": {
                               "sql": "sql", "service": "service", "group": "group", "reason": "reason", "envs": "envs",
                           }}}}
        self.assertEqual(paas.apply_configuration_errors(config), [
            "PaaS apply_contract 未配置执行窗口字段。", "PaaS 工单操作端点或字段未配置。",
        ])
        config["paas"]["task_api_base"] = "https://example.invalid"
        config["paas"]["task_actions"] = {name: {"endpoint": name, "id_field": "id"}
                                          for name in ("approve", "execute", "recall")}
        self.assertEqual(paas.apply_configuration_errors(config), ["PaaS apply_contract 未配置执行窗口字段。"])
        config["paas"]["apply_contract"]["fields"].update({"start": "start", "end": "end"})
        self.assertEqual(paas.apply_configuration_errors(config), [])

    def test_check_readiness_checks_the_requested_test_service_database(self):
        config = {"test_db": {"host": "db.example.invalid", "user": "readonly"},
                  "services": {"orders": {}}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(
                setup.check_readiness("db", service="orders"),
                ["服务未配置测试数据库：orders。"],
            )
        config["services"]["orders"]["test_database"] = "orders_test"
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_test_db_password", return_value="password"), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("db", service="orders"), [])

    def test_check_readiness_requires_the_test_database_password(self):
        config = {"test_db": {"host": "db.example.invalid", "user": "readonly"},
                  "services": {"orders": {"test_database": "orders_test"}}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(paas, "get_test_db_password", return_value=""), \
             patch.object(setup, "missing_dependencies", return_value=[]):
            self.assertEqual(setup.check_readiness("db", service="orders"), ["测试库密码未配置。"])

    def test_check_readiness_rejects_invalid_profile_before_accessing_credentials(self):
        with patch.object(paas, "load_config", side_effect=AssertionError("不应读取配置")) as load_config, \
             patch.object(paas, "get_cookie", side_effect=AssertionError("不应读取凭据")) as get_cookie, \
             patch.object(setup, "missing_dependencies", return_value=[]):
            errors = setup.check_readiness("paas", profile="bad name")
        self.assertEqual(len(errors), 1)
        self.assertIn("profile", errors[0])
        load_config.assert_not_called()
        get_cookie.assert_not_called()

    def test_setup_rejects_service_check_for_a_non_database_capability(self):
        with patch.object(sys, "argv", ["setup.py", "--check", "--capability", "sls", "--service", "orders"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            setup.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("--service 仅能与 --capability db 或 all 一起使用。", stderr.getvalue())

    def test_install_checks_requested_capability(self):
        with tempfile.TemporaryDirectory() as directory:
            bin_dir = Path(directory) / "bin"
            bin_dir.mkdir()
            capture = Path(directory) / "args.txt"
            fake_python = bin_dir / "python3"
            fake_python.write_text("#!/usr/bin/env bash\nprintf '%s\\n' \"$*\" > \"$CAPTURE_FILE\"\nexit 2\n",
                                   encoding="utf-8")
            fake_python.chmod(0o700)
            environment = {**os.environ, "PATH": str(bin_dir) + os.pathsep + os.environ["PATH"],
                           "CAPTURE_FILE": str(capture)}
            result = subprocess.run(["bash", str(Path(__file__).parent.parent / "install.sh"),
                                     "--capability", "sls"], capture_output=True, text=True, env=environment)
            self.assertEqual(result.returncode, 1)
            self.assertIn("--check-dependencies --capability sls", capture.read_text(encoding="utf-8"))

    def test_install_rejects_unknown_capability_before_checking_dependencies(self):
        result = subprocess.run(["bash", str(Path(__file__).parent.parent / "install.sh"),
                                 "--capability", "unknown"], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertIn("用法：", result.stderr)

    def test_install_checks_credentials_and_does_not_suggest_env_cookie(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / ".service-ops").mkdir()
            (root / ".service-ops" / "config.yaml").write_text("profiles: {}\n", encoding="utf-8")
            bin_dir = root / "bin"
            bin_dir.mkdir()
            capture = root / "args.txt"
            fake_python = bin_dir / "python3"
            fake_python.write_text("#!/usr/bin/env bash\nprintf '%s\\n' \"$*\" >> \"$CAPTURE_FILE\"\n", encoding="utf-8")
            fake_python.chmod(0o700)
            environment = {**os.environ, "HOME": str(root), "PATH": str(bin_dir) + os.pathsep + os.environ["PATH"],
                           "CAPTURE_FILE": str(capture)}
            result = subprocess.run(["bash", str(Path(__file__).parent.parent / "install.sh"), "--capability", "paas"],
                                    capture_output=True, text=True, errors="replace", env=environment)
            calls = capture.read_text(encoding="utf-8") if capture.exists() else ""
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("--check --capability paas", calls)
        self.assertNotIn("PAAS_COOKIE", result.stdout + result.stderr)

    def test_missing_dependencies_reports_all_runtime_modules(self):
        def find_spec(name):
            if name == "aliyun.log":
                raise ModuleNotFoundError("aliyun")
            return object()

        with patch.object(importlib.util, "find_spec", side_effect=find_spec):
            self.assertEqual(setup.missing_dependencies(), ["aliyun-log-python-sdk"])
            self.assertEqual(setup.missing_dependencies("paas"), [])
            self.assertEqual(setup.missing_dependencies("sls"), ["aliyun-log-python-sdk"])
            self.assertEqual(setup.missing_dependencies("db"), [])


class EvalContractTests(unittest.TestCase):
    def test_evaluation_rejects_non_object_result_lines(self):
        contract = {"cases": [{"id": "log-search", "expected_action": "activate",
                                "expected_route": "sls_query.py"}]}
        with self.assertRaisesRegex(ValueError, "对象"):
            eval_contract.evaluate(contract, [1])

    def test_trigger_evaluation_contract_rejects_duplicate_case_ids(self):
        contract = {"cases": [
            {"id": "same", "expected_action": "activate", "expected_route": "sls_query.py"},
            {"id": "same", "expected_action": "ask", "expected_route": "collect input"},
        ]}
        with self.assertRaisesRegex(ValueError, "重复"):
            eval_contract.evaluate(contract, [])

    def test_behavior_evaluation_contract_checks_route_questions_and_output(self):
        contract = {"cases": [{
            "id": "ambiguous", "expected_route": [], "must_ask": ["service", "sql"],
            "must_not": ["db_apply.py"], "required_output": ["下一步"],
        }]}
        result = eval_contract.evaluate(contract, [{
            "id": "ambiguous", "actual_route": [], "actual_asked": ["service"],
            "actual_output": ["下一步", "使用 db_apply.py"],
        }])
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["failed"], [{"id": "ambiguous", "reasons": ["ask", "must_not"]}])

    def test_behavior_evaluation_contract_rejects_forbidden_route_substrings(self):
        contract = {"cases": [{
            "id": "ticket", "expected_route": ["db_apply.py preflight"], "must_not": ["db_apply.py"],
        }]}
        result = eval_contract.evaluate(contract, [{
            "id": "ticket", "actual_route": ["db_apply.py preflight"], "actual_asked": [], "actual_output": [],
        }])
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["failed"], [{"id": "ticket", "reasons": ["must_not"]}])

    def test_trigger_evaluation_contract_accepts_matching_results(self):
        contract = {"cases": [{"id": "log-search", "expected_action": "activate",
                                "expected_route": "sls_query.py"}]}
        result = eval_contract.evaluate(
            contract, [{"id": "log-search", "actual_action": "activate", "actual_route": "sls_query.py"}]
        )
        self.assertEqual(result["status"], "passed")
        self.assertEqual(result["failed"], [])
        self.assertEqual(result["invalid_results"], [])

    def test_trigger_evaluation_contract_reports_route_mismatches(self):
        contract = {"cases": [{"id": "log-search", "expected_action": "activate",
                                "expected_route": "sls_query.py"}]}
        result = eval_contract.evaluate(contract, [{"id": "log-search", "actual_action": "activate",
                                                    "actual_route": "db_query.py"}])
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["failed"], [{"id": "log-search", "reasons": ["route"]}])

    def test_trigger_evaluation_contract_rejects_duplicate_results(self):
        contract = {
            "cases": [
                {"id": "log-search", "expected_action": "activate", "expected_route": "sls_query.py"}
            ]
        }

        result = eval_contract.evaluate(
            contract,
            [
                {"id": "log-search", "actual_action": "activate", "actual_route": "sls_query.py"},
                {"id": "log-search", "actual_action": "activate", "actual_route": "sls_query.py"},
            ],
        )

        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["invalid_results"], [{"id": "log-search", "reason": "duplicate"}])

    def test_trigger_evaluation_contract_rejects_unknown_and_missing_results(self):
        contract = {"cases": [{"id": "log-search", "expected_action": "activate",
                                "expected_route": "sls_query.py"}]}
        result = eval_contract.evaluate(
            contract, [{"id": "other", "actual_action": "activate", "actual_route": "sls_query.py"}]
        )
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["failed"], [{"id": "log-search", "reasons": ["missing"]}])
        self.assertEqual(result["invalid_results"], [{"id": "other", "reason": "unknown"}])


class LocalSlsQueryTests(unittest.TestCase):
    def test_sls_dry_run_never_loads_configuration_or_credentials(self):
        with patch.object(paas, "load_config", side_effect=AssertionError("不能读取配置")) as load_config, \
             patch.object(sls_query, "get_credentials", side_effect=AssertionError("不能读取凭据")) as get_credentials, \
             patch.object(sls_query, "require_sdk", side_effect=AssertionError("不能创建 SLS 客户端")) as require_sdk, \
             patch.object(sys, "argv", ["sls_query.py", "--dry-run", "--service", "orders",
                                         "--profile", "regional", "--query", "level:ERROR", "--from", "1h"]), \
             redirect_stdout(io.StringIO()) as output:
            sls_query.main()
        result = json.loads(output.getvalue())
        self.assertEqual(result["status"], "dry_run")
        self.assertEqual(result["profile"], "regional")
        self.assertEqual(result["logstore"], "orders")
        self.assertEqual(result["query"], "level:ERROR")
        self.assertEqual(result["time_range"]["duration_seconds"], 3600)
        load_config.assert_not_called()
        get_credentials.assert_not_called()
        require_sdk.assert_not_called()

    def test_sls_dry_run_supports_trace_level_and_keyword_inputs(self):
        cases = (
            (["--trace", "trace-1"], None),
            (["--level", "ERROR"], None),
            (["timeout"], '"timeout"'),
        )
        for arguments, expected_query in cases:
            with self.subTest(arguments=arguments), \
                 patch.object(paas, "load_config", side_effect=AssertionError("不能读取配置")), \
                 patch.object(sls_query, "get_credentials", side_effect=AssertionError("不能读取凭据")), \
                 patch.object(sls_query, "require_sdk", side_effect=AssertionError("不能创建 SLS 客户端")), \
                 patch.object(sys, "argv", ["sls_query.py", "--dry-run", "--service", "orders", *arguments]), \
                 redirect_stdout(io.StringIO()) as output:
                sls_query.main()
            self.assertEqual(json.loads(output.getvalue())["query"], expected_query)

    def test_sls_query_requires_an_explicit_wide_range_acknowledgement_and_caps_limit(self):
        cases = (
            (["--service", "orders", "--query", "*", "--from", "2h"], "超过 1 小时"),
            (["--service", "orders", "--query", "*", "--limit", "1001"], "1 到 1000"),
        )
        for arguments, message in cases:
            with self.subTest(arguments=arguments), \
                 patch.object(sys, "argv", ["sls_query.py", *arguments]), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                sls_query.main()
            self.assertEqual(exited.exception.code, 2)
            self.assertIn(message, stderr.getvalue())

    def test_sls_query_defaults_to_a_narrow_raw_log_window(self):
        parser = sls_query.build_parser()
        self.assertEqual(parser.parse_args([]).from_time, "15m")
        self.assertTrue(parser.parse_args([]).reverse)
        self.assertFalse(parser.parse_args(["--forward"]).reverse)

    def test_sls_query_configures_the_sdk_request_timeout(self):
        class Client:
            timeout = None

            def get_index_config(self, *_args):
                return {"keys": {}}

            def get_log_all(self, *_args, **_kwargs):
                return []

        client = Client()
        config = {"sls": {"project": "project", "region": "region"}, "log_analysis": {}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(sls_query, "require_sdk", return_value=lambda *_args: client), \
             patch.object(sys, "argv", ["sls_query.py", "--service", "orders", "--query", "*"]), \
             redirect_stdout(io.StringIO()):
            sls_query.main()
        self.assertEqual(client.timeout, sls_query.SLS_REQUEST_TIMEOUT_SECONDS)

    def test_sls_query_deadline_includes_index_discovery(self):
        class Client:
            timeout = None

            def get_index_config(self, *_args):
                return {"keys": {}}

            def get_log_all(self, *_args, **_kwargs):
                return []

        client = Client()
        config = {"sls": {"project": "project", "region": "region"}, "log_analysis": {}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(sls_query, "require_sdk", return_value=lambda *_args: client), \
             patch.object(sls_query.time, "monotonic", side_effect=(0, 0, sls_query.SLS_QUERY_DEADLINE_SECONDS + 1)), \
             patch.object(sys, "argv", ["sls_query.py", "--service", "orders", "--query", "*"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            sls_query.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("SLS 查询超过", stderr.getvalue())

    def test_sls_hard_deadline_interrupts_blocking_sdk_work(self):
        if not sls_query.supports_hard_deadline():
            self.skipTest("当前平台不支持 SLS 硬 deadline")
        with self.assertRaises(TimeoutError):
            with sls_query.hard_deadline(0.01):
                time.sleep(0.1)

    def test_sls_query_stops_paging_when_its_total_deadline_expires(self):
        page = type("Page", (), {"get_logs": lambda self: [object()]})()
        client = type("Client", (), {"get_log_all": lambda self, *_args, **_kwargs: iter([page, page])})()
        clock = iter((0, 0, sls_query.SLS_QUERY_DEADLINE_SECONDS + 1))
        with self.assertRaises(TimeoutError):
            sls_query.fetch_logs(client, "project", "orders", 0, 1, "*", 100, monotonic=lambda: next(clock))

    def test_sls_query_rejects_an_incomplete_sdk_page(self):
        page = type("Page", (), {"get_logs": lambda self: [], "is_completed": lambda self: False})()
        client = type("Client", (), {"get_log_all": lambda self, *_args, **_kwargs: iter([page])})()
        with self.assertRaisesRegex(RuntimeError, "不完整"):
            sls_query.fetch_logs(client, "project", "orders", 0, 1, "*", 100)

    def test_sls_query_reports_incomplete_pages_without_printing_partial_logs(self):
        page = type("Page", (), {"get_logs": lambda self: [],
                                  "is_completed": lambda self: False})()
        client = types.SimpleNamespace(timeout=None, get_index_config=lambda *_args: {"keys": {}},
                                       get_log_all=lambda *_args, **_kwargs: iter([page]))
        config = {"sls": {"project": "project", "region": "region"}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(sls_query, "require_sdk", return_value=lambda *_args: client), \
             patch.object(sys, "argv", ["sls_query.py", "--service", "orders"]), \
             redirect_stdout(io.StringIO()) as stdout, redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            sls_query.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertEqual(stdout.getvalue(), "")
        self.assertIn("不完整", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())

    def test_sls_sdk_requirement_is_constrained_to_the_verified_minor_version(self):
        requirements = (Path(__file__).parent.parent / "requirements.txt").read_text(encoding="utf-8")
        self.assertIn("aliyun-log-python-sdk>=0.9.50,<0.10", requirements)

    def test_sls_query_reads_all_sls_settings_from_private_config(self):
        config = {"sls": {"project": "log-project", "region": "ap-southeast-1",
                           "access_key": "config-ak", "access_secret": "config-sk"}}
        with patch.object(sls_query.paas, "_keychain_value", return_value=""), \
             patch.dict(os.environ, {"SLS_LOG_AK": "environment-ak", "SLS_LOG_SK": "environment-sk"}, clear=False):
            self.assertEqual(sls_query.get_credentials(None, config), ("config-ak", "config-sk"))
            self.assertEqual(sls_query.resolve_connection(None, None, None, None, config),
                             ("log-project", "ap-southeast-1", None))

    def test_sls_keychain_credentials_take_precedence_over_private_config(self):
        config = {"sls": {"access_key": "config-ak", "access_secret": "config-sk"}}
        with patch.object(sls_query.paas, "_keychain_value", side_effect=["keychain-ak", "keychain-sk"]):
            self.assertEqual(sls_query.get_credentials("regional", config), ("keychain-ak", "keychain-sk"))

    def test_sls_query_rejects_environment_credential_fallback(self):
        with patch.object(sls_query.paas, "_keychain_value", return_value=""), \
             patch.dict(os.environ, {"SLS_LOG_AK": "environment-ak", "SLS_LOG_SK": "environment-sk"}, clear=False):
            with self.assertRaisesRegex(ValueError, "SLS access_key/access_secret"):
                sls_query.get_credentials("default", {"sls": {}})

    def test_sls_query_uses_the_active_profile_when_no_route_is_given(self):
        with patch.object(sls_query.paas, "active_profile", return_value="regional"):
            self.assertEqual(sls_query.resolve_profile(None, None), "regional")

    def test_sls_query_requires_a_field_index_for_exact_level_filtering(self):
        with self.assertRaisesRegex(ValueError, "level.*索引"):
            sls_query.build_index_aware_query(None, "trace-1", [], "ERROR", set())
        self.assertEqual(sls_query.build_index_aware_query(None, "trace-1", [], "ERROR", {"level"}),
                         ('"trace-1" and level: ERROR', "indexed"))

    def test_sls_dry_run_marks_level_filter_as_unverified_without_index_metadata(self):
        with patch.object(sys, "argv", ["sls_query.py", "--dry-run", "--service", "orders", "--level", "ERROR"]), \
             redirect_stdout(io.StringIO()) as output:
            sls_query.main()
        result = json.loads(output.getvalue())
        self.assertEqual(result["query_mode"], "index_unverified")
        self.assertIsNone(result["query"])
        self.assertEqual(result.get("requested_filters"), {"trace": None, "terms": [], "level": "ERROR"})

    def test_sls_dry_run_marks_trace_filter_as_unverified_without_index_metadata(self):
        with patch.object(sys, "argv", ["sls_query.py", "--dry-run", "--service", "orders", "--trace", "trace-1"]), \
             redirect_stdout(io.StringIO()) as output:
            sls_query.main()
        result = json.loads(output.getvalue())
        self.assertEqual(result["query_mode"], "index_unverified")
        self.assertIsNone(result["query"])
        self.assertEqual(result.get("requested_filters"), {"trace": "trace-1", "terms": [], "level": None})

    def test_sls_detects_indexed_json_subfields(self):
        index = {"keys": {"level": {"type": "text"}, "content": {
            "type": "json", "json_keys": {"level": {"type": "text"}, "traceId": {"type": "text"}},
        }}}
        self.assertEqual(sls_query.get_indexed_fields(index), {"level", "content.level", "content.traceId", "content"})
        self.assertEqual(sls_query.build_index_aware_query(None, None, [], "ERROR", {"content.level"}),
                         ("content.level: ERROR", "indexed"))

    def test_sls_level_query_fails_closed_when_index_is_unavailable(self):
        class Client:
            timeout = None
            def get_index_config(self, *_args):
                return {"keys": {}}
            def get_log_all(self, *_args, **_kwargs):
                raise AssertionError("不应执行非精确级别查询")
        config = {"sls": {"project": "project", "region": "region"}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(sls_query, "require_sdk", return_value=lambda *_args: Client()), \
             patch.object(sys, "argv", ["sls_query.py", "--service", "orders", "--level", "ERROR"]), \
             redirect_stderr(io.StringIO()) as stderr, \
             self.assertRaises(SystemExit) as exited:
            sls_query.main()
        self.assertEqual(exited.exception.code, 2)
        self.assertIn("level 字段索引", stderr.getvalue())
        self.assertNotIn("Traceback", stderr.getvalue())

    def test_sls_level_query_uses_a_json_subfield_index(self):
        queries = []
        class Client:
            timeout = None
            def get_index_config(self, *_args):
                return {"keys": {"content": {"type": "json", "json_keys": {"level": {"type": "text"}}}}}
            def get_log_all(self, *_args, **kwargs):
                queries.append(kwargs["query"])
                return iter(())
        with patch.object(paas, "load_config", return_value={"sls": {"project": "project", "region": "region"}}), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(sls_query, "require_sdk", return_value=lambda *_args: Client()), \
             patch.object(sys, "argv", ["sls_query.py", "--service", "orders", "--level", "ERROR"]), \
             redirect_stdout(io.StringIO()):
            sls_query.main()
        self.assertEqual(queries, ["content.level: ERROR"])

    def test_sls_jsonl_keeps_authoritative_timestamp_when_content_has_time_key(self):
        content = '{"_time_":42,"level":"ERROR","message":"boom"}'
        class Log:
            def get_contents(self):
                return {"content": content}
            def get_time(self):
                return 100
        class Client:
            timeout = None
            def get_index_config(self, *_args):
                return {"keys": {}}
            def get_log_all(self, *_args, **_kwargs):
                page = type("Page", (), {"get_logs": lambda self: [Log()],
                                          "is_completed": lambda self: True})()
                return iter([page])
        config = {"sls": {"project": "project", "region": "region"}}
        with patch.object(paas, "load_config", return_value=config), \
             patch.object(sls_query, "get_credentials", return_value=("ak", "sk")), \
             patch.object(sls_query, "require_sdk", return_value=lambda *_args: Client()), \
             patch.object(sys, "argv", ["sls_query.py", "--service", "orders", "--jsonl"]), \
             redirect_stdout(io.StringIO()) as output:
            sls_query.main()
        event = json.loads(output.getvalue())
        self.assertEqual(event["_time_"], 100)
        self.assertEqual(event["message"], "boom")

    def test_parsed_json_content_does_not_leak_sensitive_fields_when_redacted(self):
        flat = sls_query.flatten_contents({"content": '{"password":"hunter2","message":"boom"}'})
        redacted = paas.redact_log_fields(flat)
        self.assertNotIn("hunter2", json.dumps(redacted))

    def test_sls_query_rejects_invalid_time(self):
        with self.assertRaises(ValueError):
            sls_query.parse_time("yesterday", 0)

    def test_sls_jsonl_aliases_keep_legacy_analysis_compatible(self):
        parser = sls_query.build_parser()
        self.assertTrue(parser.parse_args(["--raw"]).jsonl)
        self.assertTrue(parser.parse_args(["--jsonl"]).jsonl)
        self.assertTrue(parser.parse_args(["--analysis"]).jsonl)
        self.assertTrue(parser.parse_args(["--redact"]).redact)

    def test_custom_profile_names_use_one_shared_validation_rule(self):
        self.assertEqual(paas.validate_profile_name("custom-profile"), "custom-profile")
        with self.assertRaises(ValueError):
            paas.validate_profile_name("dev.us")
        with self.assertRaises(ValueError):
            paas.set_active_profile("dev.us")


class CustomProfileTests(unittest.TestCase):
    def test_database_scripts_defer_custom_profile_validation_to_config(self):
        cases = (
            (db_query.main, ["db_query.py", "--env", "prod", "--service", "wms", "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1", "--profile", "custom"]),
            (db_export.main, ["db_export.py", "--env", "prod", "--service", "wms", "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1", "--output", "/tmp/custom-profile.xlsx", "--profile", "custom"]),
            (db_apply.main, ["db_apply.py", "--action", "approve", "--id", "1", "--profile", "custom"]),
        )
        for main, argv in cases:
            with self.subTest(script=argv[0]), \
                 patch.object(paas, "set_active_profile"), \
                 patch.object(paas, "load_config", side_effect=ValueError("未配置 profile：custom。")), \
                 patch.object(sys, "argv", argv), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                main()
            self.assertEqual(exited.exception.code, 2)
            self.assertIn("未配置 profile：custom。", stderr.getvalue())

    def test_database_scripts_reject_invalid_profile_without_traceback(self):
        cases = (
            (db_query.main, ["db_query.py", "--env", "prod", "--service", "wms", "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1", "--profile", "dev.us"]),
            (db_export.main, ["db_export.py", "--env", "prod", "--service", "wms", "--purpose", "incident verification", "--sql", "SELECT id FROM orders LIMIT 1", "--output", "/tmp/invalid-profile.xlsx", "--profile", "dev.us"]),
            (db_apply.main, ["db_apply.py", "--profile", "dev.us"]),
        )
        for main, argv in cases:
            with self.subTest(script=argv[0]), \
                 patch.object(sys, "argv", argv), \
                 redirect_stderr(io.StringIO()) as stderr, \
                 self.assertRaises(SystemExit) as exited:
                main()
            self.assertEqual(exited.exception.code, 2)
            self.assertIn("profile 必须以字母开头", stderr.getvalue())
            self.assertNotIn("Traceback", stderr.getvalue())


class SlsTests(unittest.TestCase):
    def test_trace_fetch_defaults_to_a_narrow_time_window(self):
        with patch.object(sls_log_fetcher.sls_query, "resolve_profile", return_value="regional"), \
             patch.object(sls_log_fetcher, "local_sls_query_script", return_value=Path("/tools/sls_query.py")), \
             patch.object(sls_log_fetcher, "fetch_service", return_value={"status": "ok"}) as fetch_service, \
             patch.object(sys, "argv", ["sls_log_fetcher.py", "--trace-id", "trace-1", "--service", "wms"]), \
             redirect_stdout(io.StringIO()):
            sls_log_fetcher.main()
        self.assertIn("15m", fetch_service.call_args.args[0])

    def test_trigger_evaluations_have_a_valid_routing_contract(self):
        contract_path = Path(__file__).parent.parent / "evals" / "triggers.json"
        with contract_path.open(encoding="utf-8") as file:
            contract = json.load(file)
        self.assertEqual(contract["version"], 1)
        self.assertEqual(contract["skill"], "skills-service-ops")
        self.assertGreaterEqual(len(contract["cases"]), 6)
        for case in contract["cases"]:
            with self.subTest(case=case.get("id")):
                self.assertTrue(case.get("id"))
                self.assertTrue(case.get("input"))
                self.assertIn(case.get("expected_action"), {"activate", "ask", "do_not_activate"})
    def test_sls_doctor_probes_read_permission_with_one_redacted_row_limit(self):
        class Response:
            def get_logs(self):
                return [object()]

            def is_completed(self):
                return True

        calls = []
        client = types.SimpleNamespace(
            get_logstore=lambda *_args: None,
            get_index_config=lambda *_args: {"keys": {"traceId": {}}},
            get_log=lambda *args, **kwargs: calls.append((args, kwargs)) or Response(),
        )
        result = sls_query.doctor_result(client, "project", "logstore", now=1000, monotonic=lambda: 1.25)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(result["indexed_fields"], ["traceId"])
        self.assertEqual(result["read_probe"], {"line_limit": 1, "matched_rows": 1, "duration_ms": 0})
        self.assertEqual(calls, [(("project", "logstore", 700, 1000),
                                  {"query": "*", "size": 1, "offset": 0, "reverse": True})])

    def test_sls_doctor_rejects_an_incomplete_read_probe(self):
        response = type("Response", (), {"get_logs": lambda self: [],
                                         "is_completed": lambda self: False})()
        client = types.SimpleNamespace(
            get_logstore=lambda *_args: None,
            get_index_config=lambda *_args: {"keys": {}},
            get_log=lambda *_args, **_kwargs: response,
        )
        with self.assertRaisesRegex(RuntimeError, "不完整"):
            sls_query.doctor_result(client, "project", "orders", now=1000, monotonic=lambda: 1.25)

    def test_sls_doctor_uses_the_sdk_get_log_signature(self):
        response = type("Response", (), {"get_logs": lambda self: [],
                                         "is_completed": lambda self: True})()
        calls = []
        client = types.SimpleNamespace(
            get_logstore=lambda *_args: None,
            get_index_config=lambda *_args: {"keys": {}},
            get_log=lambda *args, **kwargs: calls.append((args, kwargs)) or response,
            get_logs=lambda *_args, **_kwargs: response,
        )
        result = sls_query.doctor_result(client, "project", "orders", now=1000, monotonic=lambda: 1.25)
        self.assertEqual(result["status"], "ok")
        self.assertEqual(calls, [(("project", "orders", 700, 1000),
                                  {"query": "*", "size": 1, "offset": 0, "reverse": True})])

    def test_sls_doctor_stops_when_its_total_deadline_expires(self):
        client = type("Client", (), {
            "get_logstore": lambda *_args: None,
            "get_index_config": lambda *_args: {"keys": {}},
            "get_log": lambda *_args, **_kwargs: None,
        })()
        clock = iter((0, 0, sls_query.SLS_QUERY_DEADLINE_SECONDS + 1))
        with self.assertRaises(TimeoutError):
            sls_query.doctor_result(client, "project", "logstore", now=1000,
                                    monotonic=lambda: next(clock))

    def test_trace_fetch_uses_the_bundled_query_script(self):
        self.assertEqual(sls_log_fetcher.local_sls_query_script(),
                         Path(__file__).parent / "sls_query.py")

    def test_trace_fetch_uses_raw_mode(self):
        command = sls_log_fetcher.build_sls_query_command(
            Path("/tools/sls_query.py"), "trace-1", "wms", None, "regional", "6h", "now", 100,
        )
        self.assertEqual(command, [
            sys.executable, "/tools/sls_query.py", "--trace", "trace-1", "--service", "wms",
            "--from", "6h", "--to", "now", "--limit", "100", "--raw", "--profile", "regional",
        ])

    def test_trace_fetch_forwards_explicit_redaction(self):
        command = sls_log_fetcher.build_sls_query_command(
            Path("/tools/sls_query.py"), "trace-1", "wms", None, "regional", "15m", "now", 100, redact=True,
        )
        self.assertEqual(command[-1], "--redact")

    def test_trace_fetch_normalizes_sls_log_raw_event(self):
        event = sls_log_fetcher.normalize_event({
            "_time_": 0, "level": "ERROR", "message": "boom", "traceId": "trace-1",
        })
        self.assertEqual(event["time"], "1970-01-01T00:00:00+00:00")
        self.assertEqual(event["level"], "ERROR")
        self.assertEqual(event["message"], "boom")
        self.assertEqual(event["traceId"], "trace-1")

    def test_trace_fetch_preserves_log_content_by_default(self):
        event = sls_log_fetcher.normalize_event({
            "_time_": 0, "message": "token=top-secret", "password": "hunter2", "traceId": "trace-1",
        })
        self.assertEqual(event["message"], "token=top-secret")
        self.assertNotIn("password", event)
        self.assertEqual(event["traceId"], "trace-1")

    def test_sls_query_always_marks_a_full_page_as_truncated(self):
        self.assertEqual(sls_query.truncation_metadata(100, 100, 2), {"next_page": 3, "truncated": True})
        self.assertEqual(sls_query.truncation_metadata(101, 100, 2), {"next_page": 3, "truncated": True})
        self.assertIsNone(sls_query.truncation_metadata(99, 100, 2))

    def test_sls_query_returns_next_offset_for_offset_pagination(self):
        self.assertEqual(sls_query.truncation_metadata(100, 100, offset=500),
                         {"next_offset": 600, "truncated": True})

    def test_sls_query_failure_is_reported_as_evidence_failure(self):
        result = sls_log_fetcher.fetch_service(
            ["sls-query"],
            runner=lambda *_args, **_kwargs: type("Run", (), {"returncode": 1, "stdout": "", "stderr": "failed"})(),
        )
        self.assertEqual(result["status"], "error")
        self.assertIn("--doctor", result["error"])

    def test_trace_fetch_enforces_a_subprocess_timeout(self):
        def runner(_command, **kwargs):
            self.assertEqual(kwargs["timeout"], sls_log_fetcher.SLS_QUERY_TIMEOUT_SECONDS)
            return type("Run", (), {"returncode": 1, "stdout": "", "stderr": "failed"})()

        result = sls_log_fetcher.fetch_service(["sls-query"], runner=runner)
        self.assertEqual(result["status"], "error")

    def test_trace_fetch_reports_a_subprocess_timeout(self):
        result = sls_log_fetcher.fetch_service(
            ["sls-query"],
            runner=lambda *_args, **_kwargs: (_ for _ in ()).throw(subprocess.TimeoutExpired("sls-query", 60)),
        )
        self.assertEqual(result["status"], "error")
        self.assertIn("超时", result["error"])

    def test_trace_fetch_preserves_wide_range_guidance(self):
        result = sls_log_fetcher.fetch_service(
            ["sls-query"],
            runner=lambda *_args, **_kwargs: type("Run", (), {
                "returncode": 2, "stdout": "",
                "stderr": "查询范围超过 1 小时；先 dry-run 确认范围，再传 --allow-wide-range 执行。\n",
            })(),
        )
        self.assertEqual(result["status"], "error")
        self.assertIn("--allow-wide-range", result["error"])
        self.assertNotIn("--doctor", result["error"])

    def test_trace_fetch_propagates_json_truncation_metadata(self):
        result = sls_log_fetcher.fetch_service(
            ["sls-query"],
            runner=lambda *_args, **_kwargs: type("Run", (), {
                "returncode": 0,
                "stdout": '{"_time_": 1, "level": "ERROR", "message": "boom"}\n',
                "stderr": '{"next_page": 2, "truncated": true}\n',
            })(),
        )
        self.assertTrue(result["truncated"])

    def test_profile_does_not_send_the_default_environment(self):
        commands = []
        with patch.object(sls_log_fetcher, "local_sls_query_script", return_value=Path("/tools/sls_query.py")), \
             patch.object(sls_log_fetcher, "fetch_service", side_effect=lambda command: commands.append(command) or {"status": "ok"}), \
             patch.object(sys, "argv", ["sls_log_fetcher.py", "--trace-id", "trace-1", "--service", "wms", "--profile", "custom"]), \
             redirect_stdout(io.StringIO()):
            sls_log_fetcher.main()
        self.assertIn("--profile", commands[0])
        self.assertNotIn("--env", commands[0])

    def test_trace_fetch_records_the_resolved_profile_for_analysis(self):
        with patch.object(sls_log_fetcher.sls_query, "resolve_profile", return_value="regional"), \
             patch.object(sls_log_fetcher, "local_sls_query_script", return_value=Path("/tools/sls_query.py")), \
             patch.object(sls_log_fetcher, "fetch_service", return_value={"status": "ok"}), \
             patch.object(sys, "argv", ["sls_log_fetcher.py", "--trace-id", "trace-1", "--service", "wms", "--profile", "regional"]), \
             redirect_stdout(io.StringIO()) as output:
            sls_log_fetcher.main()
        self.assertEqual(json.loads(output.getvalue())["profile"], "regional")

    def test_analyzer_sorts_timeline_and_preserves_evidence_completeness(self):
        result = log_analyzer.analyze({"services": {
            "wms": {"status": "ok", "truncated": True, "error_logs": [], "warn_logs": [
                {"time": "2026-01-01 10:02:00", "level": "WARN", "message": "late"},
                {"time": "2026-01-01 10:01:00", "level": "WARN", "message": "early"},
            ], "context_logs": []},
            "chs": {"status": "error"},
        }})
        self.assertEqual([item["message"] for item in result["timeline"]], ["early", "late"])
        self.assertEqual(result["summary"]["evidence_failed_services"], ["chs"])
        self.assertEqual(result["summary"]["truncated_services"], ["wms"])

    def test_analyzer_uses_the_configured_related_service_pattern(self):
        result = log_analyzer.analyze({"services": {"wms": {
            "status": "ok", "error_logs": [{"message": "upstream svc-pay-rpc failed"}],
            "warn_logs": [], "context_logs": [],
        }}}, related_service_pattern=r"svc-([a-z]+)-rpc")
        self.assertEqual(result["related_services"], ["pay"])

    def test_analyzer_uses_configured_entity_patterns(self):
        result = log_analyzer.analyze({"services": {"wms": {
            "status": "ok", "error_logs": [{"message": "invoice=INV-42"}],
            "warn_logs": [], "context_logs": [],
        }}}, entity_patterns={"invoice_ids": r"invoice=(INV-\d+)"})
        self.assertEqual(result["entities"], {"invoice_ids": ["INV-42"]})

    def test_analyzer_preserves_sensitive_message_content_by_default(self):
        result = log_analyzer.analyze({"services": {"wms": {
            "status": "ok", "error_logs": [{"level": "ERROR", "message": "token=top-secret"}],
            "warn_logs": [], "context_logs": [],
        }}})
        self.assertIn("top-secret", json.dumps(result))

    def test_analyzer_redacts_only_when_patterns_are_explicitly_supplied(self):
        result = log_analyzer.analyze({"services": {"wms": {
            "status": "ok", "error_logs": [{"level": "ERROR", "message": "customer=C-42"}],
            "warn_logs": [], "context_logs": [],
        }}}, redaction_patterns=[r"C-\d+"])
        self.assertNotIn("C-42", json.dumps(result))

    def test_analyzer_rejects_an_invalid_private_redaction_pattern(self):
        with self.assertRaises(ValueError):
            paas.log_analysis_rules({"log_analysis": {"redaction_patterns": ["("]}})

    def test_analyzer_rejects_malformed_collector_data(self):
        with self.assertRaises(ValueError):
            log_analyzer.analyze({"services": []})

    def test_analyzer_rejects_unknown_or_missing_service_status(self):
        for item in ({"status": "pending"}, {}):
            with self.subTest(item=item), self.assertRaisesRegex(ValueError, "status"):
                log_analyzer.analyze({"services": {"wms": item}})

    def test_analyzer_cli_works_without_private_config(self):
        with tempfile.TemporaryDirectory() as directory:
            missing_config = Path(directory) / "config.yaml"
            output = io.StringIO()
            with patch.object(log_analyzer.paas, "CONFIG_PATH", missing_config), \
                 patch.object(log_analyzer.paas, "load_config") as load_config, \
                 patch.object(sys, "stdin", io.StringIO('{"services": {}}')), \
                 redirect_stdout(output):
                log_analyzer.main([])
        self.assertEqual(json.loads(output.getvalue())["summary"]["error_count"], 0)
        load_config.assert_not_called()

    def test_analyzer_cli_uses_optional_log_analysis_config(self):
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "config.yaml"
            config_path.touch()
            output = io.StringIO()
            with patch.object(log_analyzer.paas, "CONFIG_PATH", config_path), \
                 patch.object(log_analyzer.paas, "load_config", return_value={"log_analysis": {
                     "related_service_pattern": r"svc-([a-z]+)-rpc", "entity_patterns": {},
                 }}), \
                 patch.object(sys, "stdin", io.StringIO(
                     '{"services": {"wms": {"status": "ok", "error_logs": '
                     '[{"message": "svc-pay-rpc failed"}], "warn_logs": [], "context_logs": []}}}')), \
                 redirect_stdout(output):
                log_analyzer.main([])
        self.assertEqual(json.loads(output.getvalue())["related_services"], ["pay"])

    def test_analyzer_cli_uses_the_profile_embedded_by_the_collector(self):
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "config.yaml"
            config_path.touch()
            output = io.StringIO()
            with patch.object(log_analyzer.paas, "CONFIG_PATH", config_path), \
                 patch.object(log_analyzer.paas, "load_config", return_value={"log_analysis": {}}) as load_config, \
                 patch.object(sys, "stdin", io.StringIO('{"profile": "regional", "services": {}}')), \
                 redirect_stdout(output):
                log_analyzer.main([])
        self.assertEqual(json.loads(output.getvalue())["summary"]["error_count"], 0)
        load_config.assert_called_once_with("regional")


if __name__ == "__main__":
    unittest.main()
